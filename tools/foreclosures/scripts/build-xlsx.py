#!/usr/bin/env python3
"""Build the master workbook from out/<date>/<county>.json state files.
Usage: python3 scripts/build-xlsx.py out/2026-09-08 [out/2026-10-05 ...] -o reports/foreclosures.xlsx
"""
import json, sys, glob, os, re
from datetime import datetime, date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

args=sys.argv[1:]; out='foreclosures.xlsx'
if '-o' in args: i=args.index('-o'); out=args[i+1]; del args[i:i+2]
dirs=args or sorted(glob.glob('out/*'))

def money(s):
    if s is None or s=='': return None
    m=re.search(r'-?[\d,]+(?:\.\d+)?', str(s).replace('$',''))
    return float(m.group(0).replace(',','')) if m else None
def parse_date(s):
    if not s: return None
    t=re.sub(r'^[A-Za-z]+day,?\s+','',str(s)).replace(',',', ').replace('  ',' ').strip()   # drop "Tuesday, "
    for f in ('%B %d, %Y','%m/%d/%Y','%B %d %Y','%Y-%m-%d'):
        try: return datetime.strptime(t, f).date()
        except: pass
    return None

COLS=['County','Sale date','Sale #','Case #','Status','Address','TMS','Plaintiff','Defendant','Deficiency','Judgment $','As of','Per diem $','Days to sale','Est. owed sale day $','Interest','County FMV $','Debt ÷ FMV','Acreage','Land use / type','Year built','Sq ft','Record owner','Last sale','Order / judgment','Notice','Docket','Property card','Write-up']
rows=[]
for d in dirs:
    for f in sorted(glob.glob(os.path.join(d,'*.json'))):
        st=json.load(open(f)); county=st['county']
        for r in st['rows']:
            j=r.get('judgment') or {}; p=r.get('property') or {}; n=r.get('notice') or {}; s=r.get('sources') or {}
            sale=parse_date(r.get('saleDate') or st.get('saleDate')) or parse_date(os.path.basename(d))
            asof=parse_date(j.get('asOfDate'))
            rows.append([county.title(), sale, r.get('saleNo'), r.get('caseNo'), r.get('status'), r.get('address'), r.get('tms') or n.get('tms'), r.get('plaintiff'), r.get('defendant'),
                (r.get('deficiency') or j.get('deficiency') or n.get('deficiency') or 'unknown'), j.get('totalDebt'), asof, j.get('perDiem'), None, None, j.get('interestRate') or n.get('bidInterestRate'),
                money(p.get('fmv') or p.get('assessed')), None, p.get('acreage'), p.get('landUse') or p.get('buildingType'), p.get('yearBuilt'), p.get('sqft'), p.get('owner'),
                ' '.join(x for x in [p.get('deedDate') or p.get('lastSaleDate'), p.get('lastSalePrice')] if x), j.get('source') or s.get('orderPdf'), s.get('notice') or s.get('advert') or n.get('source'), (r.get('index') or {}).get('url'), p.get('source'), r.get('writeup')])

wb=Workbook(); ws=wb.active; ws.title='All cases'
hdr=Font(name='Arial',bold=True,color='FFFFFF'); fill=PatternFill('solid',fgColor='1F3864'); body=Font(name='Arial',size=10)
ws.append(COLS)
for c in ws[1]: c.font=hdr; c.fill=fill; c.alignment=Alignment(wrap_text=True,vertical='center')
for r in rows: ws.append(r)
n=len(rows)+1
for i in range(2,n+1):
    ws[f'N{i}']=f'=IF(OR(B{i}="",L{i}=""),"",MAX(0,B{i}-L{i}))'                      # days from as-of to sale
    ws[f'O{i}']=f'=IF(K{i}="","",K{i}+IF(OR(M{i}="",N{i}=""),0,M{i}*N{i}))'         # judgment + per diem × days
    ws[f'R{i}']=f'=IF(OR(Q{i}="",Q{i}=0,O{i}=""),"",O{i}/Q{i})'                     # debt ÷ FMV
    for col in 'BL': ws[f'{col}{i}'].number_format='mm/dd/yyyy'
    for col in 'KMOQ': ws[f'{col}{i}'].number_format='$#,##0.00;($#,##0.00);-'
    ws[f'R{i}'].number_format='0%'
    for c in ws[i]: c.font=body; c.alignment=Alignment(vertical='top',wrap_text=c.column_letter in 'FHIAC')
widths={'A':11,'B':11,'C':6,'D':17,'E':10,'F':34,'G':16,'H':30,'I':28,'J':11,'K':14,'L':11,'M':10,'N':7,'O':15,'P':9,'Q':13,'R':9,'S':8,'T':18,'U':8,'V':8,'W':22,'X':18,'Y':12,'Z':12,'AA':10,'AB':12,'AC':70}
for k,v in widths.items(): ws.column_dimensions[k].width=v
ws.freeze_panes='E2'; ws.auto_filter.ref=f'A1:{get_column_letter(len(COLS))}{n}'
for col in ('Y','Z','AA','AB'):
    for i in range(2,n+1):
        c=ws[f'{col}{i}']
        if c.value and str(c.value).startswith('http'): c.hyperlink=c.value; c.value='open'; c.font=Font(name='Arial',size=10,color='0563C1',underline='single')

# Summary sheet – formulas over All cases
sm=wb.create_sheet('Summary')
sm['A1']='Master-in-Equity sales — summary'; sm['A1'].font=Font(name='Arial',bold=True,size=13)
sm['A2']='All figures are formulas over the "All cases" sheet; re-run the crawler and paste new rows there.'; sm['A2'].font=Font(name='Arial',italic=True,size=9)
head=['County','Cases','Scheduled','Cancelled / withdrawn','Deficiency demanded','Deficiency waived','With judgment $','Total judgment $ (scheduled)','With FMV']
sm.append([]); sm.append(head)
for c in sm[4]: c.font=hdr; c.fill=fill; c.alignment=Alignment(wrap_text=True)
counties=sorted({r[0] for r in rows})
for k,cty in enumerate(counties, start=5):
    sm[f'A{k}']=cty
    sm[f'B{k}']=f"=COUNTIF('All cases'!$A:$A,A{k})"
    sm[f'C{k}']=f"=COUNTIFS('All cases'!$A:$A,A{k},'All cases'!$E:$E,\"scheduled\")"
    sm[f'D{k}']=f"=B{k}-C{k}"
    sm[f'E{k}']=f"=COUNTIFS('All cases'!$A:$A,A{k},'All cases'!$J:$J,\"demanded\")"
    sm[f'F{k}']=f"=COUNTIFS('All cases'!$A:$A,A{k},'All cases'!$J:$J,\"waived\")"
    sm[f'G{k}']=f"=COUNTIFS('All cases'!$A:$A,A{k},'All cases'!$K:$K,\">0\")"
    sm[f'H{k}']=f"=SUMIFS('All cases'!$K:$K,'All cases'!$A:$A,A{k},'All cases'!$E:$E,\"scheduled\")"
    sm[f'I{k}']=f"=COUNTIFS('All cases'!$A:$A,A{k},'All cases'!$Q:$Q,\">0\")"
    sm[f'H{k}'].number_format='$#,##0;($#,##0);-'
    for c in sm[k]: c.font=body
t=5+len(counties); sm[f'A{t}']='Total'; sm[f'A{t}'].font=Font(name='Arial',bold=True)
for col in 'BCDEFGHI': sm[f'{col}{t}']=f'=SUM({col}5:{col}{t-1})'; sm[f'{col}{t}'].font=Font(name='Arial',bold=True)
sm[f'H{t}'].number_format='$#,##0;($#,##0);-'
for col,w in zip('ABCDEFGHI',[14,8,10,12,12,12,12,20,10]): sm.column_dimensions[col].width=w

# Calendar sheet
cal=wb.create_sheet('Sale calendar')
cal.append(['County','Sale day rule','Time','Location','Next dates','Notes']); [setattr(c,'font',hdr) or setattr(c,'fill',fill) for c in cal[1]]
CAL=[
 ('Greenville','1st Monday (Tue if holiday)','11:00 AM','Greenville County Courthouse, 305 E North St, judicial wing','Sep 8 · Oct 5 · Nov 2 · Dec 7 2026','5% deposit; deficiency re-bid day 30 (Oct 8 for the Sept sale)'),
 ('Pickens','1st Monday (Tue if holiday)','11:00 AM','Pickens County Courthouse, 2nd floor, Courtroom #1','Sep 8 · Oct 5 · Nov 2 · Dec 7 2026','Deficiency re-bid Oct 8 2026 11:00 AM, Courtroom #4 or #2'),
 ('Spartanburg','1st Monday (Tue if holiday)','11:00 AM','180 Magnolia St, 4th floor, Courtroom 4-A','Sep 8 · Oct 5 · Nov 2 · Dec 7 2026','Deficiency sales ~30 days later (county posts a Deficiency Sale doc)'),
 ('Oconee','1st Monday (Tue if holiday)','11:00 AM','Oconee County Courthouse, Walhalla','Sep 8 · Oct 5 · Nov 2 · Dec 7 2026','List on the SC Public Index roster (Master In Equity)'),
 ('Horry','1st Monday (Tue if holiday)','11:00 AM','1301 2nd Ave, 3rd floor, Conway','Sep 8 · Oct 5 · Nov 2 · Dec 7 2026','New bidders register 7 days ahead; first-time winners bring $2,500 certified'),
 ('Charleston','1st TUESDAY','11:00 AM','County Council Chambers, PSB, 4045 Bridge View Dr, N. Charleston','Oct 6 · Nov 3 · Dec 1 2026','REGISTER by noon the Monday before (form + photo ID); re-open sales Thursdays at 100 Broad St, Courtroom 2A'),
 ('Georgetown','1st Monday','12:00 NOON','401 Cleland St, 2nd floor MIE courtroom','Sep 8 (confirm) · Oct 5 · Nov 2 · Dec 7 2026','Holiday rule is "next day or next Monday" – confirm September (843-546-3103)'),
 ('Beaufort','1st Monday (next business day if holiday)','11:00 AM','102 Ribaut Rd, 2nd floor, Beaufort','Sep 8 · Oct 5 · Nov 2 · Dec 7 2026','List: notebook outside MIE office + Public Index roster'),
 ('Berkeley','1st WEDNESDAY','11:00 AM','Courtroom B, 300-B California Ave, Moncks Corner','Oct 7 · Nov 4 · Dec 2 2026','MIE (843) 719-4432'),
 ('Dorchester','UNVERIFIED – assumed 1st Monday','11:00 AM (assumed)','St. George courthouse','Oct 5 · Nov 2 · Dec 7 2026 (confirm)','Confirm with Clerk (843) 563-0120'),
 ('Colleton','UNVERIFIED – assumed 1st Monday','11:00 AM (assumed)','Walterboro courthouse','Oct 5 · Nov 2 · Dec 7 2026 (confirm)','Confirm with Clerk (843) 549-5791'),
]
for r in CAL: cal.append(list(r))
for col,w in zip('ABCDEF',[13,30,17,52,34,70]): cal.column_dimensions[col].width=w
for row in cal.iter_rows(min_row=2):
    for c in row: c.font=body; c.alignment=Alignment(wrap_text=True,vertical='top')

# Notes
nt=wb.create_sheet('Read me')
notes=[
 'How to read this workbook',
 'Judgment $ = the total the court entered (Order of Foreclosure / Form 4). Horry and Charleston publish it on their sale lists; other counties need the court-index run from Brian\'s PC (tools/foreclosures in the hello-world repo).',
 'Est. owed sale day = Judgment $ + per diem × days between the as-of date and the sale. This is the plaintiff\'s credit-bid ceiling, not the opening bid.',
 'Deficiency demanded = bidding stays open 30 days; anyone can top the day-of bid at the re-open. Waived = sale is final that day. "conflicting" = the notice says both; call plaintiff\'s counsel.',
 'Debt ÷ FMV under ~70% = real equity, expect competition; over 100% = plaintiff likely takes it back.',
 'County FMV / owner / acreage come from the county property card (Greenville live; qPublic counties need the PC run).',
 'Source of each row: Greenville = MIE sale list PDF (transcribed); Pickens = county roster PDF; Spartanburg = Spartan Weekly notices; Horry = county Principal Sales page; Georgetown = county DocumentCenter list; Charleston = Master\'s running list.',
 f'Generated {date.today().isoformat()} by tools/foreclosures/scripts/build-xlsx.py.',
]
for i,t_ in enumerate(notes, start=1): nt[f'A{i}']=t_; nt[f'A{i}'].font=Font(name='Arial',bold=(i==1),size=11 if i==1 else 10); nt[f'A{i}'].alignment=Alignment(wrap_text=True,vertical='top')
nt.column_dimensions['A'].width=120
wb.move_sheet('Summary', offset=-1)
wb.save(out); print('wrote',out,'rows',len(rows))
